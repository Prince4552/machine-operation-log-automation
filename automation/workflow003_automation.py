#!/usr/bin/env python3
"""WORKFLOW_003 enterprise-style expense automation prototype v2.

Input: CSV or XLSX/XLSM (all qualifying worksheets). A directory may be supplied.
Output per input: processed workbook, calculated.csv, exceptions.csv, summary.json,
state.json, audit.jsonl. A manifest.json is created for directory runs.

Calculation rules are explicit/configurable and are documented inside the output
workbook. No external accounting system is contacted; submit is a deterministic mock.
"""
from __future__ import annotations
import argparse, csv, hashlib, json, re, sys
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from statistics import median
from typing import Any, Iterable

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, Alignment
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.utils import get_column_letter
except ImportError:
    Workbook = load_workbook = None

SUPPORTED_SUFFIXES={".xlsx",".xlsm",".csv"}
REQUIRED_FIELDS=("expense_id","employee_ref","department","expense_date","category","merchant","project_code","payment_method","reimbursable","receipt_available","quantity","unit_amount","currency","fx_rate_to_reporting","tax_rate","discount_rate","source_reference")
OUTPUT_COLUMNS=("source_file","source_sheet","source_row","expense_id","employee_ref","department","expense_date","month","category","merchant","project_code","payment_method","reimbursable","receipt_available","quantity","unit_amount","currency","fx_rate_to_reporting","subtotal_amount","discount_rate","discount_amount","taxable_amount","tax_rate","tax_amount","calculated_amount","reporting_calculated_amount","duplicate_expense_id","potential_duplicate","receipt_exception","amount_exception","validation_status")
TRUE={"1","true","t","yes","y"}; FALSE={"0","false","f","no","n"}
MONEY_FIELDS=("unit_amount","subtotal_amount","discount_amount","taxable_amount","tax_amount","calculated_amount","reporting_calculated_amount")


def die(msg): raise SystemExit(f"ERROR: {msg}")
def now_utc(): return datetime.now(timezone.utc).isoformat().replace("+00:00","Z")
def q(x:Decimal,p=2): return x.quantize(Decimal("1").scaleb(-p), rounding=ROUND_HALF_UP)
def D(x,field):
    try: y=Decimal(str(x).strip())
    except (InvalidOperation,AttributeError): raise ValueError(f"{field} must be numeric")
    if not y.is_finite(): raise ValueError(f"{field} must be finite")
    return y
def parse_bool(x,field):
    s=str(x).strip().casefold()
    if s in TRUE:return True
    if s in FALSE:return False
    raise ValueError(f"{field} must be boolean-like")
def parse_date(x):
    s=str(x).strip()
    for f in ("%Y-%m-%d","%Y/%m/%d","%d/%m/%Y"):
        try:return datetime.strptime(s,f)
        except ValueError:pass
    try:return datetime.fromisoformat(s)
    except ValueError:raise ValueError(f"expense_date is not parseable: {s!r}")
def norm(x): return re.sub(r"[^a-z0-9]+","_",str(x or "").strip().casefold()).strip("_")
def canonical(obj): return json.dumps(obj,ensure_ascii=False,sort_keys=True,separators=(",",":"))
def file_sha(path): return hashlib.sha256(path.read_bytes()).hexdigest()
def write_json(path,obj): path.parent.mkdir(parents=True,exist_ok=True); path.write_text(json.dumps(obj,indent=2,ensure_ascii=False,default=str)+"\n",encoding="utf-8")
def load_json(path): return json.loads(path.read_text(encoding="utf-8"))
def audit(path,event,payload):
    with path.open("a",encoding="utf-8") as f:f.write(json.dumps({"timestamp":now_utc(),"event":event,"payload":payload},ensure_ascii=False)+"\n")


def validate_rules(r):
    need={"schema_version","reporting_currency","amount_decimal_places","supported_currencies","allowed_categories","max_line_amount_by_category","receipt_required_categories"}
    miss=sorted(need-r.keys())
    if miss: die(f"rules missing: {miss}")
    if int(r["amount_decimal_places"]) not in (0,2): die("amount_decimal_places must be 0 or 2")

def header_map(row): return {norm(v):i for i,v in enumerate(row) if norm(v)}
def detect_header(rows):
    req={norm(x) for x in REQUIRED_FIELDS}; best=None
    for i,row in enumerate(rows[:20]):
        hm=header_map(row); score=len(req & hm.keys())
        if best is None or score>best[0]:best=(score,i,hm)
    if not best or best[0] < len(req)-2: raise ValueError("could not find expense-data header")
    return best[1],best[2]


def iter_xlsx(path):
    if load_workbook is None: raise RuntimeError("openpyxl is required for Excel input/output")
    wb=load_workbook(path,read_only=True,data_only=True)
    try:
        found=False
        for ws in wb.worksheets:
            it=ws.iter_rows(values_only=True)
            probe=[]
            try:
                for _ in range(20):
                    probe.append(next(it))
            except StopIteration:
                pass
            if not probe: continue
            hi,hm=detect_header([list(r) for r in probe]); found=True
            for rn,row in enumerate(probe[hi+1:],hi+2):
                if any(v not in (None,"") for v in row):
                    yield ws.title,rn,{k:(row[pos] if pos<len(row) else None) for k,pos in hm.items()}
            start=21
            for rn,row in enumerate(it,start):
                if not any(v not in (None,"") for v in row): continue
                yield ws.title,rn,{k:(row[pos] if pos<len(row) else None) for k,pos in hm.items()}
        if not found: raise ValueError("no qualifying worksheet")
    finally: wb.close()

def iter_csv(path):
    with path.open("r",encoding="utf-8-sig",newline="") as f:
        rd=csv.DictReader(f)
        if not rd.fieldnames: raise ValueError("CSV has no header")
        for rn,row in enumerate(rd,2):yield "CSV",rn,{norm(k):v for k,v in row.items()}

@dataclass(frozen=True)
class Row:
    source_file:str; source_sheet:str; source_row:int; expense_id:str; employee_ref:str; department:str
    expense_date:datetime; category:str; merchant:str; project_code:str; payment_method:str
    reimbursable:bool; receipt_available:bool; quantity:Decimal; unit_amount:Decimal; currency:str
    fx_rate:Decimal; tax_rate:Decimal; discount_rate:Decimal; source_reference:str
    def subtotal(self,p):return q(self.quantity*self.unit_amount,p)
    def discount(self,p):return q(self.subtotal(p)*self.discount_rate,p)
    def taxable(self,p):return q(self.subtotal(p)-self.discount(p),p)
    def tax(self,p):return q(self.taxable(p)*self.tax_rate,p)
    def total(self,p):return q(self.taxable(p)+self.tax(p),p)
    def reporting_total(self,p):return q(self.total(p)*self.fx_rate,p)


def read_input(path):
    raw=[]; suffix=path.suffix.casefold()
    if suffix==".csv": raw=list(iter_csv(path))
    elif suffix in {".xlsx",".xlsm"}: raw=list(iter_xlsx(path))
    else: raise ValueError(f"unsupported file type: {path.suffix}")
    rows=[]; issues=[]
    for sheet,rn,d in raw:
        try:
            s=lambda k: "" if d.get(norm(k)) is None else str(d.get(norm(k))).strip()
            vals=dict(
                expense_id=s("expense_id"), employee_ref=s("employee_ref"), department=s("department"),
                expense_date=parse_date(d.get(norm("expense_date"))), category=s("category"), merchant=s("merchant"),
                project_code=s("project_code"), payment_method=s("payment_method"), reimbursable=parse_bool(d.get(norm("reimbursable")),"reimbursable"),
                receipt_available=parse_bool(d.get(norm("receipt_available")),"receipt_available"), quantity=D(d.get(norm("quantity")),"quantity"),
                unit_amount=D(d.get(norm("unit_amount")),"unit_amount"), currency=s("currency").upper(), fx_rate=D(d.get(norm("fx_rate_to_reporting")),"fx_rate_to_reporting"),
                tax_rate=D(d.get(norm("tax_rate")),"tax_rate"), discount_rate=D(d.get(norm("discount_rate")),"discount_rate"), source_reference=s("source_reference"))
            if not vals["expense_id"] or not vals["employee_ref"] or not vals["department"] or not vals["category"] or not vals["merchant"] or not vals["project_code"] or not vals["payment_method"] or not vals["currency"] or not vals["source_reference"]: raise ValueError("one or more required text fields are empty")
            if vals["quantity"]<=0: raise ValueError("quantity must be > 0")
            if vals["unit_amount"]<0: raise ValueError("unit_amount must be >= 0")
            if vals["fx_rate"]<=0: raise ValueError("fx_rate_to_reporting must be > 0")
            if not (Decimal(0)<=vals["tax_rate"]<=Decimal(1)): raise ValueError("tax_rate must be between 0 and 1")
            if not (Decimal(0)<=vals["discount_rate"]<Decimal(1)): raise ValueError("discount_rate must be between 0 and 1")
            rows.append(Row(path.name,sheet,rn,**vals))
        except Exception as exc:
            issues.append({"severity":"ERROR","source_sheet":sheet,"source_row":rn,"expense_id":str(d.get("expense_id") or ""),"code":"SCHEMA_VALIDATION","message":str(exc)})
    return rows,issues


def build_issues(rows,initial,rules):
    issues=list(initial); ids=Counter(r.expense_id for r in rows)
    dupe_keys=Counter((r.employee_ref,r.expense_date.date().isoformat(),r.merchant,r.total(2),r.currency) for r in rows)
    duplicate_ids={k for k,v in ids.items() if v>1}; potential={k for k,v in dupe_keys.items() if v>1}
    cats=set(rules["allowed_categories"]); cur=set(x.upper() for x in rules["supported_currencies"]); req=set(rules["receipt_required_categories"])
    limits={k:Decimal(str(v)) for k,v in rules["max_line_amount_by_category"].items()}
    for r in rows:
        total=r.total(2)
        def add(sev,code,msg):issues.append({"severity":sev,"source_sheet":r.source_sheet,"source_row":r.source_row,"expense_id":r.expense_id,"code":code,"message":msg})
        if ids[r.expense_id]>1:add("ERROR","DUPLICATE_EXPENSE_ID","expense_id appears more than once")
        if r.currency not in cur:add("ERROR","UNSUPPORTED_CURRENCY",f"{r.currency} is not configured")
        if r.category not in cats:add("ERROR","UNKNOWN_CATEGORY",f"{r.category!r} is not configured")
        if r.category in limits and total>limits[r.category]:add("WARNING","AMOUNT_LIMIT",f"{total} exceeds category limit {limits[r.category]}")
        if r.category in req and r.reimbursable and not r.receipt_available:add("WARNING","MISSING_RECEIPT","receipt required for reimbursable category")
    return issues,duplicate_ids,potential


def make_output(rows,rules,duplicate_ids,potential):
    p=int(rules["amount_decimal_places"]); req=set(rules["receipt_required_categories"]); limits={k:Decimal(str(v)) for k,v in rules["max_line_amount_by_category"].items()}; cur=set(x.upper() for x in rules["supported_currencies"]); cats=set(rules["allowed_categories"])
    out=[]
    for r in rows:
        sub=r.subtotal(p); disc=r.discount(p); taxable=r.taxable(p); tax=r.tax(p); total=r.total(p); rep=r.reporting_total(p)
        key=(r.employee_ref,r.expense_date.date().isoformat(),r.merchant,r.total(2),r.currency)
        receipt=r.category in req and r.reimbursable and not r.receipt_available
        amount=r.category in limits and total>limits[r.category]
        status="ERROR" if r.expense_id in duplicate_ids or r.currency not in cur or r.category not in cats else ("WARNING" if receipt or amount else "VALID")
        out.append({"source_file":r.source_file,"source_sheet":r.source_sheet,"source_row":r.source_row,"expense_id":r.expense_id,"employee_ref":r.employee_ref,"department":r.department,"expense_date":r.expense_date.date().isoformat(),"month":r.expense_date.strftime("%Y-%m"),"category":r.category,"merchant":r.merchant,"project_code":r.project_code,"payment_method":r.payment_method,"reimbursable":r.reimbursable,"receipt_available":r.receipt_available,"quantity":float(r.quantity),"unit_amount":float(q(r.unit_amount,p)),"currency":r.currency,"fx_rate_to_reporting":float(r.fx_rate),"subtotal_amount":float(sub),"discount_rate":float(r.discount_rate),"discount_amount":float(disc),"taxable_amount":float(taxable),"tax_rate":float(r.tax_rate),"tax_amount":float(tax),"calculated_amount":float(total),"reporting_calculated_amount":float(rep),"duplicate_expense_id":r.expense_id in duplicate_ids,"potential_duplicate":key in potential,"receipt_exception":receipt,"amount_exception":amount,"validation_status":status})
    return out


def aggregate(out,issues,rules):
    valid=[r for r in out if r["validation_status"]!="ERROR"]
    vals=[Decimal(str(r["reporting_calculated_amount"])) for r in valid]
    def sm_reporting(field):
        return float(q(sum((Decimal(str(r[field]))*Decimal(str(r["fx_rate_to_reporting"])) for r in valid),Decimal(0)),2))
    codes=Counter(i["code"] for i in issues)
    metrics={
        "input_rows":len(out),
        "valid_rows":len(valid),
        "error_rows":sum(r["validation_status"]=="ERROR" for r in out),
        "warning_rows":sum(r["validation_status"]=="WARNING" for r in out),
        "unique_expense_ids":len({r["expense_id"] for r in out}),
        "duplicate_expense_ids":len({r["expense_id"] for r in out if r["duplicate_expense_id"]}),
        "unique_employees":len({r["employee_ref"] for r in out}),
        "unique_departments":len({r["department"] for r in out}),
        "unique_projects":len({r["project_code"] for r in out}),
        "unique_merchants":len({r["merchant"] for r in out}),
        "subtotal_reporting":sm_reporting("subtotal_amount"),
        "discount_reporting":sm_reporting("discount_amount"),
        "tax_reporting":sm_reporting("tax_amount"),
        "total_expense_reporting":float(q(sum(vals,Decimal(0)),2)) if vals else 0.0,
        "reimbursable_expense_reporting":float(q(sum((Decimal(str(r["reporting_calculated_amount"])) for r in valid if r["reimbursable"]),Decimal(0)),2)),
        "non_reimbursable_expense_reporting":float(q(sum((Decimal(str(r["reporting_calculated_amount"])) for r in valid if not r["reimbursable"]),Decimal(0)),2)),
        "average_transaction_reporting":float(q(sum(vals,Decimal(0))/Decimal(len(vals)),2)) if vals else 0.0,
        "median_transaction_reporting":float(q(Decimal(str(median(vals))),2)) if vals else 0.0,
        "min_transaction_reporting":float(min(vals)) if vals else 0.0,
        "max_transaction_reporting":float(max(vals)) if vals else 0.0,
        "receipt_missing_count":sum(not r["receipt_available"] for r in out),
        "receipt_exception_count":codes.get("MISSING_RECEIPT",0),
        "high_amount_exception_count":codes.get("AMOUNT_LIMIT",0),
        "potential_duplicate_count":sum(r["potential_duplicate"] for r in out),
    }
    def groups(k,limit=None):
        g=defaultdict(lambda:[0,Decimal(0)])
        for r in valid:g[r[k]][0]+=1;g[r[k]][1]+=Decimal(str(r["reporting_calculated_amount"]))
        arr=[{k:key,"count":v[0],"reporting_total":float(q(v[1],2))} for key,v in g.items()];arr.sort(key=lambda x:(-x["reporting_total"],str(x[k])));return arr[:limit] if limit else arr
    cur=defaultdict(lambda:[0,Decimal(0),Decimal(0)])
    for r in valid:
        key=r["currency"];cur[key][0]+=1;cur[key][1]+=Decimal(str(r["calculated_amount"]));cur[key][2]+=Decimal(str(r["reporting_calculated_amount"]))
    by_currency=[{"currency":k,"count":v[0],"local_total":float(q(v[1],2)),"reporting_total":float(q(v[2],2))} for k,v in cur.items()]
    by_currency.sort(key=lambda x:(-x["reporting_total"],x["currency"]))
    return {"reporting_currency":rules["reporting_currency"],"summary_metrics":metrics,"exception_counts":dict(codes),"receipt_coverage":round(sum(r["receipt_available"] for r in out)/len(out),4) if out else 0,"group_metrics":{"by_employee":groups("employee_ref",100),"by_department":groups("department"),"by_category":groups("category"),"by_project":groups("project_code",100),"by_merchant":groups("merchant",100),"by_payment_method":groups("payment_method"),"by_currency":by_currency,"by_month":groups("month")}}


def write_csv(path,rows,fields):
    with path.open("w",encoding="utf-8-sig",newline="") as f:
        w=csv.DictWriter(f,fieldnames=fields);w.writeheader();w.writerows(rows)


def workbook(path,out,issues,summary,rules,batch_id):
    if Workbook is None:die("openpyxl is required")
    wb=Workbook(); ws=wb.active; ws.title="Summary"; ws.append(["WORKFLOW_003 — Expense Automation"]); ws["A1"].font=Font(bold=True,size=16); ws.append(["Batch ID",batch_id]); ws.append(["Generated UTC",now_utc()]); ws.append(["Reporting currency",summary["reporting_currency"]]); ws.append([]); ws.append(["Metric","Value"])
    for c in ws[6]:c.font=Font(bold=True)
    for k,v in summary["summary_metrics"].items():ws.append([k,v])
    e=wb.create_sheet("Exceptions"); eheads=["severity","source_sheet","source_row","expense_id","code","message"];e.append(eheads)
    for c in e[1]:c.font=Font(bold=True)
    for i in issues:e.append([i.get(h,"") for h in eheads])
    d=wb.create_sheet("Calculated_Data");d.append(list(OUTPUT_COLUMNS))
    for c in d[1]:c.font=Font(bold=True)
    for r in out:d.append([r.get(k) for k in OUTPUT_COLUMNS])
    for name,rows in summary["group_metrics"].items():
        s=wb.create_sheet(name[:31]);
        if rows:
            heads=list(rows[0]);s.append(heads)
            for c in s[1]:c.font=Font(bold=True)
            for r in rows:s.append([r[h] for h in heads])
        else:s.append(["No data"])
    cr=wb.create_sheet("Calculation_Rules")
    rules_rows=[("Field","Rule"),("subtotal_amount","quantity * unit_amount"),("discount_amount","subtotal_amount * discount_rate"),("taxable_amount","subtotal_amount - discount_amount"),("tax_amount","taxable_amount * tax_rate"),("calculated_amount","taxable_amount + tax_amount"),("reporting_calculated_amount","calculated_amount * fx_rate_to_reporting"),("Control","Errors block approval/submission; warnings require review.")]
    for row in rules_rows:cr.append(row)
    for c in cr[1]:c.font=Font(bold=True)
    for s in wb.worksheets:
        s.freeze_panes="A2"; s.sheet_view.showGridLines=False
        for col in s.columns:
            cells=list(col)[:250]
            if not cells:continue
            ml=max(len(str(c.value)) if c.value is not None else 0 for c in cells)
            s.column_dimensions[get_column_letter(cells[0].column)].width=min(max(ml+2,12),42)
    hi={n:i+1 for i,n in enumerate(OUTPUT_COLUMNS)}
    for r in range(2,d.max_row+1):
        d.cell(r,hi["expense_date"]).number_format="yyyy-mm-dd"
        for k in MONEY_FIELDS:d.cell(r,hi[k]).number_format='#,##0.00'
        for k in ("tax_rate","discount_rate"):d.cell(r,hi[k]).number_format='0.00%'
    if d.max_row>=2:
        ref=f"A1:{get_column_letter(d.max_column)}{d.max_row}";t=Table(displayName="CalculatedExpenses",ref=ref);t.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showRowStripes=True,showFirstColumn=False,showLastColumn=False,showColumnStripes=False);d.add_table(t)
    wb.save(path)


def batch_id(input_path,rules_path):return "BATCH-"+hashlib.sha256((file_sha(input_path)+file_sha(rules_path)).encode()).hexdigest()[:16].upper()

def process_file(path,outroot,rules_path):
    rules=load_json(rules_path);validate_rules(rules);bid=batch_id(path,rules_path); bdir=outroot/path.stem/bid;bdir.mkdir(parents=True,exist_ok=True); sp=bdir/"state.json"
    if sp.exists():
        st=load_json(sp)
        if st.get("input_sha256")==file_sha(path) and st.get("rules_sha256")==file_sha(rules_path):return st
    ap=bdir/"audit.jsonl";audit(ap,"RUN_STARTED",{"batch_id":bid,"input":path.name})
    rows,parse_issues=read_input(path);issues,dup,potential=build_issues(rows,parse_issues,rules);out=make_output(rows,rules,dup,potential);summary=aggregate(out,issues,rules);errs=sum(i["severity"]=="ERROR" for i in issues);warns=sum(i["severity"]=="WARNING" for i in issues);state={"batch_id":bid,"status":"BLOCKED" if errs else "READY_FOR_APPROVAL","created_at":now_utc(),"input_file":str(path),"input_sha256":file_sha(path),"rules_sha256":file_sha(rules_path),"row_count":len(out),"error_count":errs,"warning_count":warns,"output_workbook":str(bdir/f"processed_{path.stem}.xlsx"),"summary":summary,"approval":None,"submission":None}
    workbook(Path(state["output_workbook"]),out,issues,summary,rules,bid);write_csv(bdir/"calculated.csv",out,list(OUTPUT_COLUMNS));write_csv(bdir/"exceptions.csv",issues,["severity","source_sheet","source_row","expense_id","code","message"]);write_json(bdir/"summary.json",summary);write_json(sp,state);audit(ap,"VALIDATED",{"errors":errs,"warnings":warns});audit(ap,"RUN_BLOCKED" if errs else "READY_FOR_APPROVAL",{"batch_id":bid});return state

def process_inputs(inputs,outroot,rules):
    files=[p for p in sorted(inputs) if p.is_file() and p.suffix.casefold() in SUPPORTED_SUFFIXES]
    if not files:die("no supported input files found")
    done=[];fails=[]
    for p in files:
        try:
            s=process_file(p,outroot,rules);done.append({"input":str(p),"batch_id":s["batch_id"],"status":s["status"],"rows":s["row_count"],"errors":s["error_count"],"warnings":s["warning_count"],"batch_dir":str(Path(s["output_workbook"]).parent)})
        except Exception as exc:fails.append({"input":str(p),"error":str(exc)})
    m={"generated_at":now_utc(),"processed_files":done,"failures":fails};write_json(outroot/"manifest.json",m);return m

def approve(bdir,approver,reason):
    s=load_json(bdir/"state.json")
    if s["status"]=="SUBMITTED":return s
    if s["status"]!="READY_FOR_APPROVAL":die(f"cannot approve state={s['status']}")
    s["status"]="APPROVED";s["approval"]={"approver":approver,"reason":reason,"timestamp":now_utc()};write_json(bdir/"state.json",s);audit(bdir/"audit.jsonl","APPROVED",s["approval"]);return s

def submit_mock(bdir):
    s=load_json(bdir/"state.json")
    if s["status"]=="SUBMITTED":return s
    if s["status"]!="APPROVED":die(f"batch must be APPROVED; current={s['status']}")
    tx="MOCK-TXN-"+hashlib.sha256(f"{s['batch_id']}|{s['input_sha256']}".encode()).hexdigest()[:16].upper();s["status"]="SUBMITTED";s["submission"]={"mode":"MOCK","transaction_id":tx,"submitted_at":now_utc(),"message":"No external system was contacted."};write_json(bdir/"state.json",s);write_json(bdir/"mock_submission.json",s["submission"]);audit(bdir/"audit.jsonl","SUBMITTED_MOCK",s["submission"]);return s

def self_test():
    import tempfile
    with tempfile.TemporaryDirectory() as td:
        r=Path(td); rules=r/"rules.json";rules.write_text(json.dumps({"schema_version":"t","reporting_currency":"JPY","amount_decimal_places":2,"supported_currencies":["JPY"],"allowed_categories":["travel","meal","other"],"max_line_amount_by_category":{"travel":"1000000","meal":"1000000","other":"1000000"},"receipt_required_categories":["travel","meal"]}),encoding="utf-8")
        data=r/"in.xlsx";wb=Workbook();ws=wb.active;ws.title="A";ws.append(REQUIRED_FIELDS)
        for i in range(1,101):ws.append([f"EXP-{i}",f"EMP-{i%5}","Finance","2026-09-19","travel" if i%2 else "meal","M","PRJ","Card",True,True,2,1000,"JPY",1,.1,0,f"S{i}"])
        ws2=wb.create_sheet("B");ws2.append(REQUIRED_FIELDS);ws2.append(["EXP-101","EMP-9","HR","2026-09-20","other","N","PRJ2","Card",False,True,1,500,"JPY",1,0,0,"S101"]);wb.save(data)
        m=process_inputs([data],r/"out",rules);assert len(m["processed_files"])==1;s=load_json(Path(m["processed_files"][0]["batch_dir"])/"state.json");assert s["status"]=="READY_FOR_APPROVAL";assert s["row_count"]==101;assert abs(s["summary"]["summary_metrics"]["total_expense_reporting"]-220500)<.01;assert (Path(m["processed_files"][0]["batch_dir"])/"processed_in.xlsx").exists()
        b=Path(m["processed_files"][0]["batch_dir"]);approve(b,"tester","reviewed");a=submit_mock(b);assert a["status"]=="SUBMITTED";tx=a["submission"]["transaction_id"];assert submit_mock(b)["submission"]["transaction_id"]==tx
        print("Preflight WORKFLOW_003 automation v2 self-test: PASS (7 assertions)")

def main():
    p=argparse.ArgumentParser();sp=p.add_subparsers(dest="cmd",required=True);sp.add_parser("self-test")
    q=sp.add_parser("process");g=q.add_mutually_exclusive_group(required=True);g.add_argument("--input");g.add_argument("--input-dir");q.add_argument("--rules",required=True);q.add_argument("--output-dir",required=True)
    a=sp.add_parser("approve");a.add_argument("--batch-dir",required=True);a.add_argument("--approver",required=True);a.add_argument("--reason",required=True)
    s=sp.add_parser("submit");s.add_argument("--batch-dir",required=True)
    st=sp.add_parser("status");st.add_argument("--batch-dir",required=True)
    args=p.parse_args()
    if args.cmd=="self-test":self_test();return 0
    if args.cmd=="process":m=process_inputs([Path(args.input)] if args.input else list(Path(args.input_dir).iterdir()),Path(args.output_dir),Path(args.rules));print(json.dumps(m,indent=2,ensure_ascii=False));return 0 if not m["failures"] else 1
    if args.cmd=="approve":print(json.dumps(approve(Path(args.batch_dir),args.approver,args.reason),indent=2,ensure_ascii=False));return 0
    if args.cmd=="submit":print(json.dumps(submit_mock(Path(args.batch_dir)),indent=2,ensure_ascii=False));return 0
    if args.cmd=="status":print(json.dumps(load_json(Path(args.batch_dir)/"state.json"),indent=2,ensure_ascii=False));return 0
    return 0

if __name__=="__main__":sys.exit(main())

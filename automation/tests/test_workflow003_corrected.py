import json
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from workflow003_automation import (  # noqa: E402
    REQUIRED_FIELDS,
    Workbook,
    load_json,
    process_inputs,
    approve,
    submit_mock,
)


class Workflow003Tests(unittest.TestCase):
    def make_rules(self, root: Path) -> Path:
        p = root / "rules.json"
        p.write_text(json.dumps({
            "schema_version": "test",
            "reporting_currency": "JPY",
            "amount_decimal_places": 2,
            "supported_currencies": ["JPY", "USD"],
            "allowed_categories": ["travel", "meal", "other"],
            "max_line_amount_by_category": {
                "travel": "1000000",
                "meal": "1000000",
                "other": "1000000",
            },
            "receipt_required_categories": ["travel", "meal"],
        }), encoding="utf-8")
        return p

    def make_book(self, root: Path) -> Path:
        p = root / "input.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "A"
        ws.append(REQUIRED_FIELDS)
        ws.append([
            "EXP-1", "EMP-1", "Finance", "2026-09-19", "travel", "M", "P",
            "Card", True, True, 2, 1000, "JPY", 1, 0.10, 0, "S1"
        ])
        ws2 = wb.create_sheet("B")
        ws2.append(REQUIRED_FIELDS)
        ws2.append([
            "EXP-2", "EMP-2", "HR", "2026-09-20", "meal", "N", "P2",
            "Card", True, True, 1, 500, "JPY", 1, 0.10, 0, "S2"
        ])
        wb.save(p)
        return p

    def test_golden_calculation_and_multisheet(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            rules = self.make_rules(root)
            data = self.make_book(root)

            manifest = process_inputs([data], root / "out", rules)
            self.assertEqual(len(manifest["processed_files"]), 1)

            state = load_json(
                Path(manifest["processed_files"][0]["batch_dir"]) / "state.json"
            )
            self.assertEqual(state["status"], "READY_FOR_APPROVAL")
            self.assertEqual(state["row_count"], 2)

            # Independent expected value:
            # (2*1000)*1.10 + (1*500)*1.10 = 2750
            self.assertEqual(
                state["summary"]["summary_metrics"]["total_expense_reporting"],
                2750.0,
            )

            wb = openpyxl.load_workbook(
                state["output_workbook"], read_only=True, data_only=True
            )
            self.assertEqual(wb["Calculated_Data"].max_row, 3)
            self.assertEqual(wb["Calculated_Data"].max_column, 31)
            self.assertIn("by_employee", wb.sheetnames)
            wb.close()

    def test_directory_batch(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            rules = self.make_rules(root)
            input_dir = root / "input"
            input_dir.mkdir()

            for i in range(3):
                p = input_dir / f"f{i}.csv"
                p.write_text(
                    "expense_id,employee_ref,department,expense_date,category,"
                    "merchant,project_code,payment_method,reimbursable,"
                    "receipt_available,quantity,unit_amount,currency,"
                    "fx_rate_to_reporting,tax_rate,discount_rate,source_reference\n"
                    f"E{i},EMP{i},Finance,2026-09-19,other,M,P,Card,"
                    f"true,true,1,100,JPY,1,0,0,S{i}\n",
                    encoding="utf-8",
                )

            manifest = process_inputs(
                list(input_dir.iterdir()), root / "out", rules
            )
            self.assertEqual(len(manifest["processed_files"]), 3)
            self.assertFalse(manifest["failures"])

    def test_duplicate_and_unsupported_currency_block(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            rules = self.make_rules(root)
            p = root / "bad.csv"

            header = ",".join(REQUIRED_FIELDS)
            row1 = (
                "E1,EMP1,Finance,2026-09-19,travel,M,P,Card,"
                "true,true,1,100,JPY,1,0.1,0,S1"
            )
            row2 = (
                "E1,EMP1,Finance,2026-09-19,travel,M,P,Card,"
                "true,true,1,100,GBP,190,0.1,0,S2"
            )
            p.write_text(
                f"{header}\n{row1}\n{row2}\n", encoding="utf-8"
            )

            manifest = process_inputs([p], root / "out", rules)
            state = load_json(
                Path(manifest["processed_files"][0]["batch_dir"]) / "state.json"
            )

            self.assertEqual(state["status"], "BLOCKED")
            self.assertGreaterEqual(state["error_count"], 2)

    def test_warning_does_not_block_clean_submission_path(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            rules = self.make_rules(root)
            p = root / "warning.csv"

            p.write_text(
                ",".join(REQUIRED_FIELDS) + "\n"
                "E1,EMP1,Finance,2026-09-19,meal,M,P,Card,"
                "true,false,1,100,JPY,1,0.1,0,S1\n",
                encoding="utf-8",
            )

            manifest = process_inputs([p], root / "out", rules)
            state = load_json(
                Path(manifest["processed_files"][0]["batch_dir"]) / "state.json"
            )

            self.assertEqual(state["status"], "READY_FOR_APPROVAL")
            self.assertGreaterEqual(state["warning_count"], 1)

    def test_idempotent_mock_submit(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            rules = self.make_rules(root)
            data = self.make_book(root)

            manifest = process_inputs([data], root / "out", rules)
            batch = Path(manifest["processed_files"][0]["batch_dir"])

            approve(batch, "tester", "reviewed")
            first = submit_mock(batch)
            second = submit_mock(batch)

            self.assertEqual(
                first["submission"]["transaction_id"],
                second["submission"]["transaction_id"],
            )
            self.assertEqual(second["status"], "SUBMITTED")


if __name__ == "__main__":
    unittest.main(verbosity=2)
